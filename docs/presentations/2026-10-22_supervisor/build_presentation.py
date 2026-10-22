from pathlib import Path
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from PIL import Image
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import json, shutil
from docx import Document
from docx.shared import Pt as DPt
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.colors import HexColor
from xml.sax.saxutils import escape

OUT=Path(__file__).resolve().parent; AS=OUT/'assets'; AS.mkdir(exist_ok=True)
ROOT=OUT.parents[1]/'thesis'; IMG=ROOT/'figures/img'
NAVY='142D45'; TEAL='087F8C'; BLUE='537FA7'; GOLD='C48835'; INK='20374B'; MUTED='526779'; BG='F6F8FA'; WHITE='FFFFFF'; LIGHT='E6EFF3'
p=Presentation();p.slide_width= Inches(13.333333);p.slide_height=Inches(7.5)
p.core_properties.title='Multi-Person Human Mesh Recovery with the Anny Body Model'
p.core_properties.author='Muhammad Saad Najib'
p.core_properties.subject='Supervisor presentation: student model, distillation, data conversion and verified results'
records=[]
def rgb(s):return RGBColor.from_string(s)
def rect(s,x,y,w,h,c,line=None,r=False):
 sh=s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE if r else MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(c);sh.line.fill.background() if not line else None
 if line:sh.line.color.rgb=rgb(line)
 if r:
  try:sh.adjustments[0]=0.12
  except:pass
 return sh

def text(s,t,x,y,w,h,size=22,color=INK,bold=False,align=None):
 sh=s.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h));tf=sh.text_frame;tf.clear();tf.word_wrap=True
 tf.margin_left=tf.margin_right= Inches(.025);tf.margin_top=tf.margin_bottom= Inches(.025)
 for i,line in enumerate(t.split('\n')):
  q=tf.paragraphs[0] if i==0 else tf.add_paragraph();q.text=line;q.font.name='Aptos';q.font.size=Pt(size);q.font.bold=bold;q.font.color.rgb=rgb(color);q.space_after=Pt(7)
  if align is not None:q.alignment=align
 return sh

def slide(title,section,source='Thesis revision, 22 Oct 2026',subtitle=None):
 s=p.slides.add_slide(p.slide_layouts[6]);s.background.fill.solid();s.background.fill.fore_color.rgb=rgb(BG)
 rect(s,0,0,13.333,.10,TEAL);text(s,section.upper(),.55,.30,11,.3,11,TEAL,True)
 text(s,title,.55,.85,12.2,.85,30,NAVY,True)
 if subtitle:text(s,subtitle,.58,1.65,12,.6,16,MUTED)
 rect(s,.55,7.0,12.2,.012,'CED9E1');text(s,source,.58,7.12,11.6,.22,9,MUTED);text(s,f'{len(p.slides):02}',12.25,7.10,.5,.25,11,TEAL,True,PP_ALIGN.RIGHT)
 records.append({'slide':len(p.slides),'title':title,'section':section,'source':source})
 return s

def notes(s,script,cue='',transition=''):
 rec=records[-1];rec.update(script=script,cue=cue,transition=transition)
 s.notes_slide.notes_text_frame.text=f'SLIDE {rec["slide"]}: {rec["title"]}\n\nSPEAKING SCRIPT\n{script}\n\nPOINT / DEMO CUE\n{cue}\n\nTRANSITION\n{transition}\n\nSOURCE\n{rec["source"]}'

def card(s,x,y,w,h,title,body,accent=TEAL):
 rect(s,x,y,w,h,WHITE,r=True);rect(s,x,y,.055,h,accent);text(s,title,x+.22,y+.18,w-.42,.65,20,accent,True);text(s,body,x+.22,y+.98,w-.42,h-1.08,18,INK)

def arrow(s,x1,y1,x2,y2,c=TEAL):
 sh=s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2));sh.line.color.rgb=rgb(c);sh.line.width=Pt(2)
 # OOXML arrowhead
 from pptx.oxml.xmlchemy import OxmlElement
 end=OxmlElement('a:tailEnd');end.set('type','triangle');sh.line._get_or_add_ln().append(end)
 return sh

def box(s,label,x,y,w,h,c=LIGHT,fs=19):
 rect(s,x,y,w,h,c,r=True);return text(s,label,x+.10,y+.16,w-.2,h-.2,fs,INK,True,PP_ALIGN.CENTER)

def pic(s,path,x,y,w,h):
 path=Path(path);im=Image.open(path);iw,ih=im.size;scale=min(w/iw,h/ih);ww=iw*scale;hh=ih*scale
 return s.shapes.add_picture(str(path),Inches(x+(w-ww)/2),Inches(y+(h-hh)/2),width=Inches(ww),height=Inches(hh))

def table(s,headers,rows,x,y,w,h,widths=None,fs=18):
 sh=s.shapes.add_table(len(rows)+1,len(headers),Inches(x),Inches(y),Inches(w),Inches(h));tb=sh.table
 if widths:
  for col,wi in zip(tb.columns,widths):col.width=Inches(wi)
 for r,vals in enumerate([headers]+rows):
  for c,v in enumerate(vals):
   cell=tb.cell(r,c);cell.text=str(v);cell.margin_left=Inches(.13);cell.margin_right=Inches(.08);cell.margin_top=Inches(.10);cell.vertical_anchor=MSO_ANCHOR.MIDDLE
   cell.fill.solid();cell.fill.fore_color.rgb=rgb(NAVY if r==0 else (WHITE if r%2 else LIGHT))
   for q in cell.text_frame.paragraphs:q.font.name='Aptos';q.font.size=Pt(fs);q.font.color.rgb=rgb(WHITE if r==0 else INK);q.font.bold=(r==0)
 return sh

def banner(s,t,y=6.15,color=TEAL):
 rect(s,.58,y,12.15,.62,color,r=True);text(s,t,.80,y+.09,11.7,.48,19,WHITE,True)

def equation(s,latex,name,x,y,w,h):
 fig=plt.figure(figsize=(14,1.4),dpi=180);fig.patch.set_alpha(0);fig.text(.5,.5,latex,ha='center',va='center',fontsize=25,color='#'+NAVY)
 file=AS/(name+'.png');fig.savefig(file,transparent=True,bbox_inches='tight',pad_inches=.12);plt.close(fig);pic(s,file,x,y,w,h)

# 1
s=slide('Multi-Person Human Mesh Recovery','01 / Introduction','Muhammad Saad Najib · Supervisor presentation · 22 October 2026')
text(s,'with the Anny Body Model',.60,1.62,11.8,.6,30,TEAL,True)
text(s,'A smaller student model for recovering 3D people\nfrom a single image',.62,2.55,6.5,1.3,27,NAVY)
text(s,'Muhammad Saad Najib\nMaster’s thesis · Supervisor presentation',.62,4.70,6.3,1.0,20,MUTED)
# real source illustration, labelled as teacher
pic(s,IMG/'final_model_real_photos.jpg',7.8,2.4,4.6,3.8)
text(s,'Illustration: Base/v5 and teacher overlays\nThese are not student predictions.',7.9,6.25,4.5,.55,12,MUTED)
notes(s,'My thesis studies how to recover a three-dimensional body for each person visible in one photograph. A body mesh is a surface made from many connected points, like a digital mannequin. It represents body pose and shape, rather than just a few joints. The main model in this presentation is my smaller student model. It learns from labelled examples and from a larger teacher. I will explain the starting system, what the student learns, how I converted real-image labels into the Anny body format, and the measured trade-off between speed and reconstruction quality. The illustrations on this first slide show earlier Base and teacher outputs; the student results will be identified separately.','Introduce the task before using abbreviations.','First, why is this problem difficult?')
# 2
s=slide('One photograph, several difficult decisions','02 / Problem statement')
card(s,.6,2.05,3.9,3.45,'Recover a whole body','Estimate pose, body shape\nand distance from the camera.\n\nPeople can overlap or be\npartly hidden.')
card(s,4.72,2.05,3.9,3.45,'Real photographs','Synthetic images have exact\nlabels but differ from photos.\n\n3DPW annotations use\na different body model.')
card(s,8.84,2.05,3.9,3.45,'Reduce inference cost','Large models require more\ncomputation and memory.\n\nCan a smaller model retain\nuseful reconstruction quality?')
banner(s,'Research question: how much accuracy can a smaller student retain?')
notes(s,'A photograph shows a flat image, but our output must describe a three-dimensional body. Several different bodies and distances can produce a similar picture, especially when clothing or another person hides part of the body. There are two additional challenges in this project. First, synthetic training images come with detailed labels, while real photos look different and have less complete supervision. Second, the larger model is expensive to run. My question is therefore whether a smaller student can learn enough from a stronger teacher and from both datasets to provide useful results with lower inference cost. I evaluate this as a trade-off: a faster result is valuable, but I also need to show how much reconstruction accuracy is lost.','Point to the three challenges from left to right.','The starting point is Multi-HMR.')
# 3
s=slide('Baseline: one image → multiple 3D bodies','03 / Original Multi-HMR','Sources: Multi-HMR (Baradel et al., 2024); released NAVER Anny checkpoint; thesis §4.2',subtitle='Baseline used in this project: the released Multi-HMR model with Anny outputs.')
labels=['RGB image\n672 × 672','DINOv2\nViT-L encoder','Detect people\n+ image offsets','Person decoder\n(cross-attention)','Anny body\n+ placement']
for i,l in enumerate(labels):
 box(s,l,.62+i*2.55,2.85,2.22,1.25,fs=18)
 if i<4:arrow(s,2.84+i*2.55,3.47,3.13+i*2.55,3.47)
text(s,'Camera intrinsics K → geometric placement',7.30,2.10,5.2,.45,18,TEAL,True)
arrow(s,11.95,2.62,11.95,2.84)
text(s,'Encoder: extracts image information',.72,4.60,5.9,.6,21)
text(s,'Decoder: predicts one body per detected person',6.50,4.60,6,.8,21)
banner(s,'Single-shot: process the whole image in one forward pass.')
notes(s,'This is the baseline pipeline. The RGB image is resized and padded to the model input size. The DINOv2 encoder turns it into image features, which can be thought of as descriptions of small image regions. A detection head identifies person anchors and their image offsets. The person decoder uses cross-attention to combine information about each person with the rest of the image. It predicts the parameters needed by a body model, and the body layer converts those parameters into a mesh. Camera information helps place the mesh in three-dimensional space. The original Multi-HMR work also describes an SMPL-X version; our external reference is specifically NAVER’s released Anny checkpoint. This distinction matters when comparing outputs and model sizes.','Trace the arrows. Define encoder and decoder in plain words.','My student preserves this general pipeline but uses a smaller encoder.')
# 4
s=slide('The student learns from labels and a frozen teacher','04 / Student–teacher architecture','Sources: thesis §4.1, §4.8; model.py; distill.py; train.py')
box(s,'Same RGB image\n+ camera K',.62,3.0,1.9,1.2,fs=18)
box(s,'Teacher: ViT-L\n1,024 features',3.0,2.05,2.35,1.0,fs=18)
box(s,'Teacher heads\n→ body parameters',6.1,2.05,2.55,1.0,fs=18)
box(s,'Anny layer\n→ mesh + joints',9.6,2.05,2.85,1.0,fs=18)
box(s,'Student: ViT-B\n768 features',3.0,4.10,2.35,1.0,'D5EEE9',18)
box(s,'Projector\n768 → 1,024',5.72,4.10,2.0,1.0,'D5EEE9',18)
box(s,'Copied heads\n→ parameters',8.08,4.10,2.05,1.0,'D5EEE9',18)
box(s,'Anny layer\n→ mesh',10.5,4.10,2.0,1.0,'D5EEE9',18)
for a,b,c,d in [(2.52,3.3,3,2.55),(2.52,3.8,3,4.6),(5.35,2.55,6.1,2.55),(8.65,2.55,9.6,2.55),(5.35,4.6,5.72,4.6),(7.72,4.6,8.08,4.6),(10.13,4.6,10.5,4.6)]:arrow(s,a,b,c,d)
text(s,'FROZEN: no weight updates',3.03,1.67,5,.3,12,MUTED,True)
arrow(s,5.1,3.06,6.55,4.08,GOLD);text(s,'Feature loss',4.0,3.35,2.1,.4,15,GOLD,True)
arrow(s,8.25,3.06,9.0,4.08,GOLD);text(s,'Output loss',8.7,3.35,2.0,.4,15,GOLD,True)
text(s,'Dataset labels → supervised losses on student predictions',3.0,5.50,9.4,.5,21,TEAL,True)
banner(s,'Deployment uses only the student branch, including its projector.')
notes(s,'Both branches receive the same image and camera information. The larger teacher has a ViT-L encoder and remains frozen, meaning training does not update its weights. The student uses a smaller ViT-B encoder. Because its feature width is 768 rather than 1,024, a small trainable projector converts its features to the width expected by the transferred prediction heads. There are three sources of learning. Dataset labels provide the supervised loss. Teacher features provide the feature-distillation loss. Teacher body parameters provide the output-distillation loss. These losses update the student, while teacher targets are detached from gradient updates. The copied heads are frozen initially, then trained. At deployment the teacher is removed, but the projector stays because the student heads need its output.','Point to gold arrows as teacher guidance and the bottom line as label supervision.','Next, what exactly does the network predict?')
# 5
s=slide('What goes in, and what comes out?','04 / Model inputs and outputs','Sources: thesis §4.2; model.py',subtitle='A body model turns predicted parameters into a visible 3D surface.')
card(s,.62,2.35,3.62,3.42,'Input','RGB image + camera K\n\nTraining also supplies\nper-person body labels\nand anchor locations.')
card(s,4.53,2.35,4.05,3.42,'Predictions per person','Detection score + 2D offset\n163 bone rotations\n11 body-shape parameters\nHead-anchor depth')
card(s,8.87,2.35,3.85,3.42,'Decoded output','3D body mesh + joints\n\nPose: how the body bends\nShape: body proportions\nPlacement: where it stands')
notes(s,'The input used at inference is an RGB image and camera intrinsics. When the camera calibration is unavailable, the demo uses an assumed field of view, which can affect recovered depth. For each detected person, the model predicts bone rotations, eleven Anny body-shape parameters, and quantities describing the head anchor’s image location and depth. Rotations describe articulation, while shape parameters control aspects such as height and body proportions. The Anny layer converts these parameters into a mesh and joints. The mesh is then translated so its head anchor reaches the predicted camera-space position. During training, labels provide target parameters and geometry. At inference those targets are absent: the model must estimate everything from the image and camera information.','Use pose, shape and placement as three separate ideas.','The losses tell the student which parts of its answer need to improve.')
# 6
s=slide('Supervised losses: compare predictions with labels','04 / Loss functions','Sources: thesis §4.3; Table T02; exact formulas in backup slides')
table(s,['Loss group','What is compared?','Plain-language purpose'],[
['Detection + offset','Person-anchor score and 2D location','Find each person correctly'],
['Pose + shape','Bone rotations and body parameters','Recover articulation and proportions'],
['Depth + translation','Head depth and 3D anchor position','Put the body at the right distance'],
['3D joints + vertices','Pelvis-centred predicted / target bodies','Match the body geometry'],
['2D projection','Predicted / target points in the image','Make the mesh line up with the photo']],.62,2.05,12.08,3.65,[2.65,5.05,4.38],17)
banner(s,'Ltotal = Lsupervised + Lfeature KD + Loutput KD')
notes(s,'A loss is a numerical penalty for a mismatch during training. Here the penalties measure several different parts of the prediction. Detection and offset losses help find people and locate their anchors. Pose and shape losses compare the predicted body parameters with their labels. Depth and translation losses teach camera-space placement. The three-dimensional joint and vertex losses compare bodies after moving each pelvis to the origin, so they measure body geometry without absolute translation. The two-dimensional losses project the body back onto the image and measure alignment there. This is why a low centred body error does not automatically mean correct placement in the photo. The student receives all these supervised terms plus the two teacher-guidance terms on the next slide.','Read purposes, not every formula. Full weights and equations are in the backup.','Distillation adds information from the teacher.')
# 7
s=slide('Knowledge distillation: two kinds of teacher guidance','05 / Training process','Sources: thesis §4.8; distill.py; loss.py')
card(s,.62,2.05,5.85,3.45,'Feature distillation','Compare teacher image features with\nthe student’s projected features.\n\nKL divergence, temperature T = 4.\n“Learn similar image representations.”')
card(s,6.75,2.05,5.95,3.45,'Output distillation','Compare corresponding people’s\nrotations, shape parameters and depth.\n\nWeighted L1 differences.\n“Learn similar body predictions.”')
banner(s,'The student learns from the true labels and the teacher; both KD weights are 1.0.')
notes(s,'Knowledge distillation means using a larger trained model to guide a smaller one. The first form is feature distillation: the student’s projected image features are encouraged to resemble the teacher’s features. The implementation turns feature channels into softened distributions and compares them with KL divergence. A temperature of four controls that softening. The second form is output distillation: the student is encouraged to predict similar rotations, body-shape parameters and depth. During training the two models use the same labelled person locations, so outputs can be compared person by person. Teacher predictions are useful guidance, but they can also contain mistakes. Ground-truth supervision remains present. The exact reductions and weights are included in the backup slides for technical questions.','Clarify that the teacher supplies guidance; it does not create new verified ground truth.','Here is the training recipe used for the reported student.')
# 8
s=slide('Training recipe for the reported student','05 / Training process','Sources: thesis Table T11; exp8_distill_vitb.sh; job 3614191')
for x,num,label in [(0.65,'30','epochs'),(3.80,'1,000','steps per epoch'),(6.95,'8','images per batch'),(10.10,'70 / 30','synthetic / real mix')]:
 text(s,num,x,2.25,2.6,.7,37,TEAL,True);text(s,label,x,3.0,2.6,.7,18,MUTED)
box(s,'Initialize ViT-B\n+ copy teacher heads',.65,4.05,3.5,1.0,fs=18);box(s,'First 5 epochs\nkeep copied heads frozen',4.90,4.05,3.5,1.0,fs=18);box(s,'Continue training\nthen calibrate student outputs',9.15,4.05,3.5,1.0,fs=18)
arrow(s,4.15,4.55,4.9,4.55);arrow(s,8.4,4.55,9.15,4.55)
text(s,'Learning rate: 3 × 10⁻⁵; halved every 12 epochs. Teacher stays frozen.',.65,5.55,12,.5,21)
text(s,'70/30 is the sampling ratio of images, not people. Test images are excluded.',.65,6.28,12,.45,17,MUTED)
notes(s,'The reported student is the completed thirty-epoch ViT-B run. Each epoch has one thousand optimizer steps, with eight images in a batch, giving two hundred and forty thousand sampled image presentations over training. These are not necessarily unique images. The sampler draws about seventy percent synthetic AnnyOne images and thirty percent converted 3DPW training images. This is an image ratio; the person ratio differs because images contain different numbers of people. The student starts with an HMR-trained ViT-B backbone and copied teacher heads. The heads stay frozen for the first five epochs, then train with the student. The learning rate starts at three times ten to the minus five and halves every twelve epochs. Student-specific output calibration is estimated after training.','Emphasize this is the confirmed 30-epoch checkpoint, not an unverified longer run.','I use several metrics because a single number does not describe the entire result.')
# 9
s=slide('How reconstruction quality is measured','05 / Evaluation metrics','Sources: thesis §4.10; Table T07',subtitle='Errors are measured on matched people. Lower error is better; higher recall is better.')
table(s,['Metric','Simple meaning','Unit'],[
['PVE','Average surface-point error after pelvis alignment','mm ↓'],
['PA-PVE','Surface error after optimal rotation, translation and scale alignment','mm ↓'],
['MPJPE','Average joint-position error after pelvis alignment','mm ↓'],
['Recall','Share of labelled people recovered by the detector','% ↑'],
['Latency / model cost','Time per forward pass; parameters and operations','ms / M / GFLOPs ↓']],.62,2.35,12.08,3.65,[2.15,8.0,1.93],17)
notes(s,'PVE is the average distance between corresponding mesh surface points after aligning the pelvis. A value of one hundred millimetres means an average point error of ten centimetres under that evaluation alignment. PA-PVE goes further: it optimally rotates, translates and scales the predicted body before measuring error. It is therefore useful for examining pose and shape, but can hide errors in global size and orientation. MPJPE performs a similar comparison using joint positions rather than surface vertices. Recall records how many labelled people were recovered, because a geometry score on matched people cannot describe missed detections. Finally, parameters, operation count and measured latency describe computational cost. These metrics answer different questions and should be presented together.','Explain millimetres with a centimetre example. Do not call PA-PVE a placement metric.','These are the measured reconstruction results.')
# 10
s=slide('The student loses accuracy relative to its teacher','05 / Reconstruction results','Sources: thesis Tables T16–T17; jobs 3614191, 3611336, 3614192',subtitle='PVE / PA-PVE, in mm. All 3DPW values use converted Anny labels.')
table(s,['Model','Synthetic holdout','3DPW test','3DPW recall'],[
['Released baseline','60.3 / 44.6*','126.9 / 68.8','97%'],
['Teacher (Final, ViT-L)','73.7 / 50.3','89.0 / 57.6','99%'],
['Student (ViT-B)','119.6 / 80.8','110.1 / 71.8','99%']],.62,2.35,12.08,2.55,[4.1,2.8,2.8,2.38],20)
text(s,'* Baseline synthetic images are potentially in its training set.\nBaseline real-image training exposure also differs from teacher/student.',.65,5.15,12,.70,16,MUTED)
banner(s,'Student: lower 3DPW PVE than the released baseline, but higher error than the teacher.')
notes(s,'The teacher is the strongest of our two trained models on these reconstruction measurements. On the synthetic holdout, the student’s PVE is 119.6 millimetres compared with 73.7 for the teacher. On the converted 3DPW test set, the student reaches 110.1 millimetres compared with 89.0 for the teacher. Both recover ninety-nine percent of labelled people under this evaluation. The student has lower 3DPW PVE than the released baseline, but its PA-PVE is slightly higher than that baseline’s. There are important comparison limits: the released baseline may have seen the synthetic holdout images, and it does not have the same real-image training exposure. Also, these 3DPW numbers use Anny labels, so they should not be compared directly with published SMPL-space scores.','Read one dataset at a time. Point out the PA-PVE reversal against the baseline.','The benefit of the student appears in computational cost.')
# 11
s=slide('The student is smaller and faster','05 / Efficiency results','Source: thesis Table T16; paired L40S measurements, job 3614191',subtitle='Same benchmark: 672-pixel input, one person, batch 1, fp32 with TF32.')
table(s,['Model','Parameters','Compute','Median forward time'],[
['Teacher','318.38 M','1,927.4 GFLOPs','59.7 ms'],['Student','101.38 M','603.3 GFLOPs','28.5 ms']],.62,2.25,12.08,1.85,[3.0,2.7,3.0,3.38],21)
for x,n,l in [(0.8,'3.14×','fewer parameters'),(5.0,'3.19×','less computation'),(9.1,'2.10×','faster forward pass')]:text(s,n,x,4.55,3.5,.7,36,TEAL,True);text(s,l,x,5.3,3.5,.5,19,MUTED)
text(s,'This is a model-forward benchmark, not end-to-end video FPS.',.65,6.37,12,.4,18,MUTED)
notes(s,'This slide compares the student with its teacher under the same recorded L40S benchmark. The student has about one hundred and one million parameters, compared with about three hundred and eighteen million for the teacher. It also uses roughly one third of the counted computation. The median model forward time falls from 59.7 to 28.5 milliseconds, giving a speedup of about 2.1 times. These measurements support the claim that the student reduces model cost. They do not prove that a complete video application runs at the reciprocal frame rate, because loading frames, preprocessing, rendering and encoding take additional time. The final conclusion combines this slide with the previous one: the student achieves substantial efficiency gains, with a measurable reconstruction-accuracy cost.','Use the paired measurements only; do not mix L40S and A100 timings.','Now I will explain the data used to train these models.')
# 12
s=slide('Two data sources with one common body format','06 / Datasets','Sources: thesis Chapter 3; numbers_data.tex')
card(s,.62,2.15,5.85,3.65,'AnnyOne: synthetic images','Exact Anny pose, shape and camera labels.\n\n579,860 training images\n100-image thesis holdout\n\n70% of mixed training samples')
card(s,6.75,2.15,5.95,3.65,'3DPW: real video frames','Original labels describe SMPL bodies.\nConvert them into Anny training targets.\n\nTrain 5,882 · Validation 2,622 · Test 8,190\n30% of mixed training samples')
text(s,'The 100-image synthetic holdout contains adults only; it cannot establish child performance.',.65,6.30,12,.5,17,MUTED)
notes(s,'AnnyOne provides synthetic images generated with the Anny body model. This gives detailed pose, shape and camera supervision in exactly the format predicted by the network. My training split uses 579,860 images, while the last one hundred images form the thesis holdout. That holdout contains adults, so it cannot answer whether the model works well for children. The real-image source is 3DPW, which consists of recorded video sequences with body annotations. However, those annotations use SMPL rather than Anny. I did not create 3DPW itself: my contribution here is converting its existing annotations into usable Anny targets while keeping the official training, validation and test separation. The corrected conversion yields the image counts shown on the right.','Say “converted 3DPW labels,” not “created the 3DPW dataset.”','The conversion is an optimization problem.')
# 13
s=slide('How the 3DPW–Anny labels were created','06 / Label conversion','Sources: thesis §3.3–3.5; smpl_to_anny.py, converter v2.3')
steps=[('1','Read source labels','3DPW joints, shape,\ncamera and frame ID'),('2','Fit an Anny body','Optimize root, pose,\nthen free shape values'),('3','Share body shape','One median shape\nfor each person’s sequence'),('4','Build image labels','Refit poses; transform\nto camera coordinates')]
for i,(n,tt,b) in enumerate(steps):
 x=.65+i*3.1;text(s,n,x,2.20,2.8,.65,34,TEAL,True);card(s,x,3.0,2.85,2.45,tt,b)
text(s,'Fit every third valid frame on a shared frame grid; validate and filter the completed labels.',.65,5.85,12,.75,20)
notes(s,'The converter begins with the source body joints and camera information from 3DPW. It searches for Anny parameters whose joints are close to those targets. The fitting proceeds in stages: root alignment first, then articulated pose, and then the shape parameters that can be supported by joint positions. A previous frame provides an initial guess for the next frame. After the first pass, the median fitted shape is assigned to the person across the sequence, and poses are fitted again with that fixed shape. This avoids treating the same actor as a different body in every frame. A label builder then transforms fits into camera coordinates and groups all labelled actors in each image. The shared sampling grid ensures both actors are represented on the same selected frames.','Walk through the four numbered stages.','The difficult part is that the two skeletons do not define every joint at the same place.')
# 14
s=slide('Virtual joints compare equivalent body locations','06 / Geometric correspondence','Sources: thesis §3.4 and Appendix B',subtitle='A geometric offset changes the point used for fitting, while the Anny skeleton stays intact.')
# schematic native shapes
box(s,'Anny bone head\nh',.9,2.7,2.9,1.0,fs=21);box(s,'Virtual point\np = h + Qδ',5.15,2.7,3.1,1.0,'D5EEE9',21);box(s,'SMPL joint\ntarget',9.65,2.7,2.8,1.0,fs=21)
arrow(s,3.8,3.2,5.15,3.2);arrow(s,8.26,3.2,9.65,3.2)
text(s,'Move by a fixed offset δ,\nrotated with the bone by Q',3.35,4.1,5.2,.9,21,TEAL,True)
text(s,'Minimize the gap\nto the target',9.0,4.1,3.4,.9,21,TEAL,True)
banner(s,'Joint definitions can be aligned; limb lengths and body proportions can still differ.')
notes(s,'SMPL and Anny can place a joint with a similar name at slightly different anatomical positions. If I simply compare the names, the optimizer may try to bend or stretch the body to remove a mismatch caused by the joint definition itself. A virtual joint gives the fitting procedure a more comparable point. I take a point attached to an Anny bone and add a fixed offset measured in the aligned rest skeletons. When the bone rotates, the offset rotates with it. In the equation, h is the bone head, delta is the local offset, and Q rotates that offset into the posed frame. The resulting point p is compared with the SMPL target. This does not change Anny’s actual rig or make the two anatomies identical.','The diagram is conceptual, not a measured anatomical drawing.','The completed annotations are checked independently of the fitting objective.')
# 15
s=slide('Converted labels: checked, with residual warnings','06 / Label validation','Sources: thesis §3.5, Table T05; F14_label_check_downtown_walking_00771.jpg')
for x,l in [(0.65,'Original image'),(4.72,'Source SMPL annotation'),(8.8,'Converted Anny annotation')]:text(s,l,x,2.0,3.9,.5,18,TEAL,True,PP_ALIGN.CENTER)
pic(s,IMG/'F14_label_check_downtown_walking_00771.jpg',.65,2.65,12.0,2.4)
text(s,'14.1 mm median body-joint discrepancy',.68,5.27,11.8,.55,27,TEAL,True)
text(s,'WARN: remaining limb-proportion and joint-definition differences.\nThis label-fit discrepancy is not the model’s prediction error.',.68,6.0,11.8,.75,18)
notes(s,'These pictures compare the real image, the source SMPL annotation, and the converted Anny annotation. They are label-conversion examples, not network predictions. An independent validation reads the completed labels and compares their joints and projections with the source annotations. The reported median body-joint discrepancy is 14.1 millimetres. The resulting diagnostic status is WARN, because the body models still differ in proportions and some joint definitions. Earlier label-integrity problems were corrected, but that does not guarantee every fit is perfect. It is especially important not to confuse this conversion discrepancy with the student’s prediction error or subtract it from PVE. The two quantities are measured on different objects and answer different questions.','Point to the source and converted meshes, then the discrepancy.','The next slides show existing qualitative reconstruction evidence.')
# 16
s=slide('Synthetic example: inspect placement as well as pose','07 / Pictures','Sources: thesis Figure F12; AnnyOne holdout image 579860, V1',subtitle='Historical comparison: Base/v5 versus the larger teacher (Final).')
for i,l in enumerate(['Input image','Dataset ground truth','Base / v5','Teacher / Final']):text(s,l,.65+i*3.05,2.28,2.9,.45,18,TEAL,True,PP_ALIGN.CENTER)
pic(s,IMG/'F12_holdout_579860_img_gt_base_final.jpg',.65,2.85,12.0,3.0)
text(s,'Meshes use each model’s own predicted depth. These panels are not student results.',.65,6.22,12,.55,18,MUTED)
notes(s,'This synthetic image allows a visual comparison with a known ground-truth mesh. The four columns are the input image, dataset ground truth, the earlier Base or v5 model, and the larger Final teacher. This comparison is useful for looking at apparent size and placement, not just whether the body bends plausibly. Each prediction is shown at its own predicted depth in the V1 evaluation setting. In this example, the teacher’s body appears closer to the ground-truth size and location than the Base prediction. This is a selected example rather than proof of improvement on every image. It also does not show student performance. I identify the models explicitly because the word Final in older experiment filenames refers to the larger teacher, while my deployed model is the student.','Point to ground truth first; compare the third and fourth panels.','Real-photo demonstrations require additional caution about rendering settings.')
# 17
s=slide('Real-photo examples show strengths and limits','07 / Pictures','Sources: thesis Figures F10; walking and garden examples; demo rendering settings')
text(s,'Image  |  Base/v5 overlay  |  Image  |  Teacher/Final overlay',.65,1.95,12,.5,20,TEAL,True,PP_ALIGN.CENTER)
pic(s,IMG/'F10_3dpw_walking_v5_vs_final.jpg',.65,2.60,12.0,1.85)
pic(s,IMG/'F10_photo_garden_v5_vs_final.jpg',.65,4.65,12.0,1.8)
text(s,'Legacy demos: box-based placement and rest-pose fingers; not raw depth or hand-accuracy evidence.',.65,6.60,12,.30,14,MUTED)
notes(s,'These are existing real-image demonstrations of Base and the larger teacher. The top row is a 3DPW walking scene, while the lower row is a real-photo garden scene. They show that the system can place plausible whole-body surfaces on multiple people, but individual overlays are not a quantitative evaluation. These legacy demos use bounding-box-based placement and force fingers to a rest pose. Therefore I cannot use them to prove correct raw camera depth or accurate articulated fingers. The garden example also reminds us that people of different sizes and ages remain difficult, especially when partially hidden. The purpose of this slide is to show representative output and its limits. A fair student comparison requires student outputs rendered under the same settings.','Do not describe these as newly run student examples.','The video pipeline provides the next comparison using the student itself.')
# 18
s=slide('Video comparison: student and released baseline','07 / Video demonstration','Sources: video-compare-3642415.out; thesis video evidence; demo_video.py',subtitle='Completed run: man_walking.mp4 · 356 frames · NVIDIA L40S')
text(s,'Same 356 frames · Student: 356/356 matched · Baseline: 354/356 matched',.65,2.10,12,.5,19,TEAL,True)
rect(s,.65,2.85,12.0,2.95,LIGHT,TEAL,r=True)
text(s,'INSERT YOUR COMPARISON VIDEO HERE',1.3,3.70,10.7,.6,27,TEAL,True,PP_ALIGN.CENTER)
text(s,'Input  |  Student  |  Released baseline',1.3,4.45,10.7,.5,20,MUTED,False,PP_ALIGN.CENTER)
text(s,'Inspect jitter, sudden pose changes and changing body shape.\nFrame coverage does not prove smoothness. Video FPS has not yet been verified.',.65,6.08,12,.76,18)
notes(s,'The completed walking-video run processes the same 356 frames with the student and NAVER’s released baseline. The tracker matched the primary person on all 356 student frames and on 354 baseline frames. These counts show successful execution and coverage, but they do not measure pose accuracy or smoothness. For the live demo, open the comparison video and inspect whether the mesh shakes, whether limbs suddenly change pose, and whether body shape varies between frames. The pipeline operates frame by frame; it does not train or apply a temporal reconstruction model. A dedicated placeholder is left on this slide so I can insert the comparison video before presenting. The newer FPS feature is implemented, but this successful log predates it and supplies no verified video FPS result.','Before presenting: replace the large placeholder with comparison.mp4 using Insert > Video > This Device. Test playback. If the clip is unavailable, report the log counts only.','The remaining issues define what we can honestly conclude.')
# 19
s=slide('What remains unresolved','07 / Limitations and next steps','Sources: thesis Discussion and Conclusion; video evidence')
card(s,.62,2.1,3.9,3.7,'Accuracy gap','Student errors remain higher\nthan the teacher’s.\n\nThe 95% accuracy-retention\ntarget is not met.')
card(s,4.74,2.1,3.9,3.7,'Real-image body shape','Children and body-shape\nvariation remain difficult.\n\nThe adult-only holdout\ncannot test child performance.')
card(s,8.86,2.1,3.85,3.7,'Temporal evidence','Review the actual video.\n\nMeasure synchronized FPS\nand frame-to-frame stability.')
notes(s,'There are three main limitations. First, the student does not retain the teacher’s reconstruction accuracy within the planned ninety-five-percent criterion. The compression succeeds computationally, but the accuracy target remains unmet. Second, body shape on real images is still difficult, particularly for children. Strong synthetic results cannot establish performance on real children, and our small synthetic holdout contains adults. Third, a successful video run does not establish temporal stability. We need to inspect the clip and report reproducible timing and motion-variation measurements. These are bounded next steps rather than claims that a large new training programme is already complete. The presentation therefore separates what was implemented, what was measured, and what still needs evidence.','Keep these as specific limitations, not a list of speculative fixes.','The final slide brings the contribution and the trade-off together.')
# 20
s=slide('A smaller student with a measured trade-off','07 / Conclusion')
for y,n,tt in [(2.10,'01','Converted real 3DPW annotations into checked Anny labels.'),(3.2,'02','Trained a ViT-B student with label and teacher supervision.'),(4.3,'03','Measured a 2.10× speedup, with higher reconstruction error.')]:
 text(s,n,.68,y,1,.6,29,TEAL,True);text(s,tt,1.8,y,10.5,.9,25,NAVY)
banner(s,'Main result: a clear speed–accuracy trade-off, plus a working video-comparison pipeline.')
notes(s,'The contribution is a complete and evaluated pipeline around a smaller student model. The data work makes real 3DPW supervision usable in Anny format, with validation of the converted annotations. The training work combines direct labels, transferred prediction heads and two distillation signals from a larger frozen teacher. The evaluation shows that the student reduces parameters and computation by roughly a factor of three and halves the recorded forward latency, but has higher reconstruction error than the teacher. The video-comparison implementation adds a practical way to examine deployment behaviour on identical frames. My conclusion is therefore a measured speed–accuracy trade-off. The immediate priorities are reviewing the actual walking clip, collecting its timing and stability evidence, and keeping the written conclusions aligned with those observations.','Pause for questions. Use the following slides only for technical details.','Backup: exact objectives and sources.')
# 21
s=slide('Backup: supervised objective and exact weights','Technical backup','Sources: thesis §4.3; Table T02; loss.py')
equation(s,r'$L_{sup}=10L_{det}+L_{off}+0.1L_{rot}+20L_{shape}$','sup1',.8,1.95,11.8,.6)
equation(s,r'$+5L_{depth}+5L_{trans}+100(L_{j3D}+L_{v3D})+0.1(L_{j2D}+L_{v2D})$','sup2',.8,2.65,11.8,.7)
equation(s,r'$L_{rot}=\frac{1}{N_h}\sum_i\sum_b m_{ib}w_b\|\hat R_{ib}-R_{ib}\|_1$','rot',.8,3.55,11.8,.8)
text(s,'L1 = sum of absolute differences; rotations use matrix entries.\nNₕ = supervised people; m = availability mask; w = normalized bone weight.\n3DPW hand rotations are masked; mesh losses still contain rest-pose hands.\nDepth and translation refer to the same head anchor used for placement.\nThe final recipe supervises all shape entries, including converter constants.',.85,4.65,11.8,2.0,18)
notes(s,'This backup gives the supervised weights used by the reported student. Detection has weight ten, offsets one, rotation zero point one, and shape twenty. Depth and translation each have weight five. The centred joint and vertex geometry terms each have weight one hundred, while the two-dimensional terms each have weight zero point one. These weights cannot be interpreted in isolation because the losses use different units and reductions. In the rotation expression, the mask suppresses unavailable 3DPW hand targets, and normalized per-bone weights emphasize selected regions. The mesh terms still include the rest-pose hands from converted labels. Shape supervision in the retained final configuration includes the converter’s fixed phenotype entries. These details matter when interpreting limitations or reproducing training.','Use only if asked about training loss or label masks.','The next backup gives the two distillation formulas.')
# 22
s=slide('Backup: distillation equations','Technical backup','Sources: thesis §4.8; distill.py; loss.py rotation_distill_term')
equation(s,r'$p^t=\mathrm{softmax}(Z^t/T),\quad p^s=\mathrm{softmax}(\tilde Z^s/T),\quad T=4$','kd0',.8,1.9,11.8,.6)
equation(s,r'$L_{feat}=\frac{T^2}{B}\sum_{u,n,c}p^t_{unc}\log\frac{p^t_{unc}}{p^s_{unc}}$','kd1',.8,2.7,11.8,.9)
equation(s,r'$L_{out}=\frac{1}{N_h}\sum_i[0.1\sum_b\|R^s_{ib}-R^t_{ib}\|_1+20\sum_k|\phi^s_{ik}-\phi^t_{ik}|+5|d^s_i-d^t_i|]$','kd2',.7,3.8,12.0,.85)
text(s,'B: batch size · n: image token · c: feature channel · i: person · b: bone · k: shape value\nFeature KL averages over the batch and sums tokens/channels.\nOutput KD uses matched labelled-person locations and detached teacher predictions.\nRotation KD does not use the supervised hand mask or per-bone boost.',.85,5.0,11.8,1.5,17)
notes(s,'For feature distillation, each token’s feature channels are converted into a probability distribution using temperature four. KL divergence compares the teacher distribution with the student distribution. The implementation multiplies by temperature squared and divides by batch size. It sums over tokens and channels, so describing it as an average over tokens would change the scale. Output distillation compares teacher and student rotations, phenotype values and focal-normalized log-depth for corresponding people. The weights are zero point one, twenty and five respectively. The teacher outputs are detached, so gradients update only the student. Unlike the supervised rotation loss, rotation distillation uses all bone residuals without the missing-label hand mask or the supervised per-bone weighting.','Explain indices only if needed.','The last slide lists the source papers and project evidence.')
# 23
s=slide('Selected references and evidence','Technical backup','Full bibliography and run provenance: docs/thesis/')
text(s,'Method background',.65,1.95,12,.45,23,TEAL,True)
text(s,'Baradel et al. (2024). Multi-HMR: Multi-Person Whole-Body Human Mesh Recovery in a Single Shot.\nBrégier et al. Human Mesh Modeling for Anny Body. arXiv:2511.03589.\nOquab et al. (2023). DINOv2: Learning Robust Visual Features without Supervision.\nvon Marcard et al. (2018). Recovering Accurate 3D Human Pose in the Wild Using IMUs and a Moving Camera.\nHinton et al. (2015). Distilling the Knowledge in a Neural Network.',.65,2.55,12.0,2.6,17)
text(s,'Project evidence',.65,5.35,12,.45,23,TEAL,True)
text(s,'Revised thesis: commit 3e2f21e · Student results: job 3614191\nVideo execution: job 3642415 · Every pictured model is identified on its slide.',.65,5.95,12,.75,18)
notes(s,'These sources provide the architecture, body model, feature encoder, dataset and distillation background. The numerical results in this presentation come from the revised repository thesis and its traceable experiment records, rather than from the sample presentation or invented estimates. The completed student evaluation is job 3614191, and the walking-video execution is job 3642415. The accompanying source map identifies the thesis sections and image files used on each slide. Existing pictures preserve the model and rendering-setting distinctions described in the thesis. Before a final thesis defence, bibliographic metadata should be checked against the official publications, as the thesis already records that author-review step.','Keep as a reference slide for questions, not a slide to read aloud.','')

# Save deck and script outputs
ppt=OUT/'Najib_supervisor_presentation.pptx';p.save(ppt)
(OUT/'slide_content.json').write_text(json.dumps(records,indent=2,ensure_ascii=False))
md=['# Speaker script — Muhammad Saad Najib','', '20 main slides + 3 technical backup slides. Suggested speaking time: about 20–25 minutes, plus video and questions.','', 'Model names: Student = deployed ViT-B; Teacher/Final = larger ViT-L; Base/v5 = earlier thesis model; released baseline = NAVER Multi-HMR-Anny.','']
doc=Document();doc.add_heading('Speaker script',0);doc.add_paragraph('Muhammad Saad Najib • 22 October 2026');doc.add_paragraph('20 main slides + 3 backup slides. About 20–25 minutes, plus the video and questions. Full scripts are also in PowerPoint speaker notes.')
styles=getSampleStyleSheet();styles.add(ParagraphStyle(name='ScriptBody',fontName='Helvetica',fontSize=11,leading=16,spaceAfter=12));styles['Heading1'].textColor=HexColor('#'+NAVY)
story=[Paragraph('Speaker script',styles['Title']),Paragraph('Muhammad Saad Najib — 22 October 2026',styles['Normal']),Spacer(1,18),Paragraph('20 main slides + 3 technical backup slides. About 20–25 minutes, plus the video and questions. Full scripts are also included in the PowerPoint speaker notes.',styles['ScriptBody'])]
for rec in records:
 md += [f'## Slide {rec["slide"]} — {rec["title"]}','',rec['script'],'',f'**Point / demo cue:** {rec["cue"]}','',f'**Transition:** {rec["transition"]}','',f'**Source:** {rec["source"]}','']
 doc.add_page_break();doc.add_heading(f'Slide {rec["slide"]}: {rec["title"]}',1);doc.add_paragraph(rec['script']);doc.add_heading('Point / demo cue',2);doc.add_paragraph(rec['cue']);doc.add_heading('Transition',2);doc.add_paragraph(rec['transition']);doc.add_paragraph('Source: '+rec['source'])
 story += [PageBreak(),Paragraph(f'Slide {rec["slide"]}: {escape(rec["title"])}',styles['Heading1']),Spacer(1,16),Paragraph(escape(rec['script']),styles['ScriptBody']),Paragraph('Point / demo cue',styles['Heading2']),Paragraph(escape(rec['cue']),styles['ScriptBody']),Paragraph('Transition',styles['Heading2']),Paragraph(escape(rec['transition']),styles['ScriptBody']),Paragraph('Source: '+escape(rec['source']),styles['Normal'])]
(OUT/'Speaker_script.md').write_text('\n'.join(line.rstrip() for line in md));doc.save(OUT/'Speaker_script.docx');SimpleDocTemplate(str(OUT/'Speaker_script.pdf')).build(story)
(OUT/'Sources_and_presenting_notes.md').write_text('''# Presentation notes\n\nBuilt from the seven topic slides in the supplied template. Dense architecture, training and dataset topics have been expanded into 20 main slides plus 3 backup slides. The supplied template contained no pictures or speaker notes.\n\n## Before presenting\n- Open the PPTX in PowerPoint and use Presenter View for the full speaking script.\n- Copy the real video from `/netscratch/najib/multi-hmr/results/video_compare/video_3642415/comparison.mp4` onto your presentation computer. Slide 18 has a dedicated video placeholder, as requested. This delivery does not contain the video.\n- On slide 18, remove the placeholder and use Insert > Video > This Device. Test video playback before the meeting.\n- The older completed video job has no synchronized FPS measurement. Do not quote its coverage counts as an accuracy or temporal-stability result.\n- Existing picture comparisons show Base/v5 versus Teacher/Final. They must not be described as student predictions. Real-photo legacy demos use box placement and rest fingers.\n- The source PowerPoint is editable: text, tables, diagrams and connectors are native objects; equations and experimental pictures are raster images.\n\n## Source map\n'''+''.join(f'- Slide {r["slide"]}: {r["source"]}\n' for r in records))
print('Saved',ppt,'slides',len(p.slides),'script words',sum(len(r['script'].split()) for r in records))
